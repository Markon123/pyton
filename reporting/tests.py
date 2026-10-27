from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from federation.models import Profile, Discipline
from .models import ReportJob
from .services import process_reports


class ReportWorkflowTests(TestCase):
    def setUp(self):
        self.folder = TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.settings_override = override_settings(LOCAL_DIR=Path(self.folder.name))
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.user = get_user_model().objects.create_user('reporter')
        Profile.objects.create(user=self.user, roles=['federation'])
        self.client.force_login(self.user)
        year = timezone.localdate().year
        self.params = {'year_from': year-1, 'year_to': year, 'format': 'xlsx'}

    def test_request_worker_and_real_workbook(self):
        response = self.client.post(reverse('reporting:create'), self.params)
        self.assertEqual(response.status_code, 302)
        job = ReportJob.objects.get()
        self.assertEqual(process_reports(), 1)
        job.refresh_from_db()
        self.assertEqual(job.status, 'ready')
        download = self.client.get(reverse('reporting:download', args=[job.number]))
        workbook = load_workbook(BytesIO(b''.join(download.streaming_content)))
        self.assertEqual(workbook.active.cell(6, 1).value, self.params['year_from'])
        self.assertEqual(process_reports(), 0)

    def test_invalid_format_and_dates_do_not_queue(self):
        for fields in ({**self.params, 'format': 'exe'}, {**self.params, 'year_from': 2099}):
            self.assertEqual(self.client.post(reverse('reporting:create'), fields).status_code, 400)
        self.assertFalse(ReportJob.objects.exists())

    def test_other_federation_account_cannot_fetch_job(self):
        self.client.post(reverse('reporting:create'), self.params)
        job = ReportJob.objects.get()
        other = get_user_model().objects.create_user('other-reporter')
        Profile.objects.create(user=other, roles=['federation'])
        self.client.force_login(other)
        self.assertEqual(self.client.get(reverse('reporting:job', args=[job.number])).status_code, 404)
        self.assertEqual(self.client.get(reverse('reporting:download', args=[job.number])).status_code, 404)

    def test_failure_does_not_publish_file(self):
        self.client.post(reverse('reporting:create'), self.params)
        with patch('reporting.services.build_census', side_effect=ValueError('sensitive details')):
            self.assertEqual(process_reports(), 0)
        job = ReportJob.objects.get()
        self.assertEqual(job.status, 'failed')
        self.assertNotIn('sensitive', job.error)
        self.assertEqual(self.client.get(reverse('reporting:download', args=[job.number])).status_code, 404)

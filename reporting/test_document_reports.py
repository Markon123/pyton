from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import uuid

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from openpyxl import load_workbook

from federation.models import (
    Application, Competition, Discipline, JudgeAssignment, Organization,
    Profile, RankApplication, Result, Team,
)
from .forms import validate_saved_params
from .models import ReportJob
from .services import expected_filename, process_reports, report_directory


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class DocumentReportWorkflowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.today = timezone.localdate()
        cls.discipline = Discipline.objects.create(name='Тестовая дисциплина', code='report-test')
        cls.organization = Organization.objects.create(name='Вымышленный клуб очереди', municipality='Оренбург')
        cls.fed = cls.person('report_fed', ['federation'], 'Федерация Очереди')
        cls.other_fed = cls.person('report_other_fed', ['federation'], 'Другая Федерация')
        cls.athlete = cls.person('report_athlete', ['athlete'], 'Вымышленный Спортсмен')
        cls.judge = cls.person('report_judge', ['judge'], 'Вымышленный Судья')
        cls.competition = Competition.objects.create(
            title='Демонстрационное соревнование очереди', starts_on=cls.today, ends_on=cls.today,
            registration_opens=cls.today - timedelta(days=7), registration_closes=cls.today,
            location='Учебный стенд', state='finished', protocol_published=True, plan_number='ДЕМО-ОЧЕРЕДЬ-1',
        )
        cls.competition.disciplines.add(cls.discipline)
        cls.application = Application.objects.create(
            competition=cls.competition, athlete=cls.athlete, discipline=cls.discipline,
            status='approved', submitted_by=cls.athlete.user,
        )
        cls.result = Result.objects.create(application=cls.application, place=1, score=Decimal('142.50'))
        JudgeAssignment.objects.create(competition=cls.competition, judge=cls.judge, position='Главный судья')
        RankApplication.objects.create(
            athlete=cls.judge, kind='category', requested_title='Тестовая судейская категория',
            submitted_by=cls.fed.user, status='approved', assigned_on=cls.today - timedelta(days=30),
            valid_until=cls.today + timedelta(days=30), order_number='ДЕМО-СУДЬЯ-1',
        )
        cls.team = Team.objects.create(
            season=cls.today.year, discipline=cls.discipline, age_min=12, age_max=99, created_by=cls.fed.user,
        )
        cls.team.members.add(cls.athlete)

    @classmethod
    def person(cls, username, roles, name):
        first, last = name.split(' ', 1)
        user = get_user_model().objects.create_user(username, first_name=first, last_name=last)
        profile = Profile.objects.create(
            user=user, roles=roles, birth_date=date(cls.today.year - 20, 1, 1),
            organization=cls.organization, start_year=cls.today.year - 3, municipality='Оренбург',
        )
        profile.disciplines.add(cls.discipline)
        return profile

    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        settings = override_settings(LOCAL_DIR=Path(temporary.name))
        settings.enable()
        self.addCleanup(settings.disable)
        self.client.force_login(self.fed.user)

    def queue(self, kind, object_id, fmt='xlsx'):
        response = self.client.post(reverse('reporting:create'), {'kind': kind, 'object_id': object_id, 'format': fmt})
        self.assertEqual(response.status_code, 302)
        return ReportJob.objects.latest('pk')

    def downloaded(self, job):
        response = self.client.get(reverse('reporting:download', args=[job.number]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Cache-Control'], 'private, no-store')
        return b''.join(response.streaming_content)

    def test_all_document_kinds_generate_real_files_in_each_format(self):
        sources = [
            ('protocol', self.competition.pk, self.competition.title, self.athlete.display_name),
            ('athlete', self.athlete.pk, self.competition.title, self.athlete.display_name),
            ('team', self.team.pk, 'Сборная региона', self.athlete.display_name),
            ('judging-panel', self.competition.pk, self.competition.title, self.judge.display_name),
            ('judging-history', self.judge.pk, self.competition.title, self.judge.display_name),
        ]
        for kind, pk, expected_title, expected_person in sources:
            for fmt in ('xlsx', 'docx', 'pdf'):
                with self.subTest(kind=kind, fmt=fmt):
                    job = self.queue(kind, pk, fmt)
                    self.assertEqual(process_reports(), 1)
                    job.refresh_from_db()
                    self.assertEqual(job.status, 'ready')
                    content = self.downloaded(job)
                    if fmt == 'xlsx':
                        workbook = load_workbook(BytesIO(content))
                        text = '\n'.join(str(value) for row in workbook.active.iter_rows(values_only=True) for value in row if value is not None)
                    elif fmt == 'docx':
                        document = Document(BytesIO(content))
                        text = '\n'.join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])
                    else:
                        self.assertTrue(content.startswith(b'%PDF-'))
                        self.assertIn(b'/FontFile2', content)
                        continue
                    self.assertIn(expected_title, text)
                    self.assertIn(expected_person, text)
        self.assertEqual(process_reports(), 0)

    def test_link_prefills_descriptive_picker_without_queueing(self):
        response = self.client.get(reverse('reporting:index'), {'kind': 'athlete', 'object_id': self.athlete.pk, 'fmt': 'pdf'})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.athlete.display_name)
        self.assertContains(response, self.organization.name)
        self.assertEqual(response.context['form'].initial['format'], 'pdf')
        self.assertEqual(str(response.context['form'].initial['object_id']), str(self.athlete.pk))
        self.assertFalse(ReportJob.objects.exists())

    def test_unknown_kinds_invalid_source_ids_and_formats_are_rejected(self):
        for values, status in [
            ({'kind': 'arbitrary', 'object_id': self.athlete.pk, 'format': 'xlsx'}, 404),
            ({'kind': 'athlete', 'object_id': -1, 'format': 'xlsx'}, 400),
            ({'kind': 'athlete', 'object_id': '1.5', 'format': 'xlsx'}, 400),
            ({'kind': 'athlete', 'object_id': self.judge.pk, 'format': 'xlsx'}, 400),
            ({'kind': 'athlete', 'object_id': self.athlete.pk, 'format': 'exe'}, 400),
        ]:
            with self.subTest(values=values):
                self.assertEqual(self.client.post(reverse('reporting:create'), values).status_code, status)
        self.assertFalse(ReportJob.objects.exists())

    def test_complete_preliminary_protocol_does_not_publish_competition(self):
        self.competition.protocol_published = False
        self.competition.state = 'closed'
        self.competition.save(update_fields=['protocol_published', 'state'])
        job = self.queue('protocol', self.competition.pk)
        self.assertEqual(process_reports(), 1)
        workbook = load_workbook(BytesIO(self.downloaded(job)))
        text = '\n'.join(str(value) for row in workbook.active.iter_rows(values_only=True) for value in row if value is not None)
        self.assertIn('Предварительный', text)
        self.assertIn(self.athlete.display_name, text)
        self.competition.refresh_from_db()
        self.assertFalse(self.competition.protocol_published)
        self.assertEqual(self.competition.state, 'closed')

    def test_incomplete_or_draft_protocol_cannot_be_queued(self):
        self.competition.protocol_published = False
        self.competition.state = 'draft'
        self.competition.save(update_fields=['protocol_published', 'state'])
        payload = {'kind': 'protocol', 'object_id': self.competition.pk, 'format': 'pdf'}
        self.assertEqual(self.client.post(reverse('reporting:create'), payload).status_code, 400)
        self.competition.state = 'closed'
        self.competition.save(update_fields=['state'])
        self.result.delete()
        self.assertEqual(self.client.post(reverse('reporting:create'), payload).status_code, 400)
        self.assertFalse(ReportJob.objects.exists())

    def test_role_revocation_prevents_generation_of_old_jobs(self):
        job = self.queue('athlete', self.athlete.pk)
        self.fed.roles = ['athlete']
        self.fed.save(update_fields=['roles'])
        self.assertEqual(process_reports(), 0)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertFalse(job.filename)
        self.assertEqual(self.client.get(reverse('reporting:job', args=[job.number])).status_code, 403)

    def test_disabled_account_cannot_generate_queued_file(self):
        job = self.queue('athlete', self.athlete.pk)
        self.fed.user.is_active = False
        self.fed.user.save(update_fields=['is_active'])
        self.assertEqual(process_reports(), 0)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')

    def test_other_account_and_non_federation_role_cannot_read_ready_file(self):
        job = self.queue('athlete', self.athlete.pk)
        self.assertEqual(process_reports(), 1)
        self.client.force_login(self.other_fed.user)
        self.assertEqual(self.client.get(reverse('reporting:job', args=[job.number]), HTTP_ACCEPT='application/json').status_code, 404)
        self.assertEqual(self.client.get(reverse('reporting:download', args=[job.number])).status_code, 404)
        self.client.force_login(self.athlete.user)
        self.assertEqual(self.client.get(reverse('reporting:download', args=[job.number])).status_code, 403)
        self.client.force_login(self.fed.user)
        self.fed.roles = ['athlete']
        self.fed.save(update_fields=['roles'])
        self.assertEqual(self.client.get(reverse('reporting:download', args=[job.number])).status_code, 403)

    def test_tampered_saved_parameters_never_generate_file(self):
        invalid = [
            {'kind': 'athlete', 'object_id': self.athlete.pk, 'path': '/etc/passwd'},
            {'kind': '../../escape', 'object_id': self.athlete.pk},
            {'kind': ['athlete'], 'object_id': self.athlete.pk},
            {'kind': 'athlete', 'object_id': True},
            {'kind': 'athlete', 'object_id': 1.5},
        ]
        for params in invalid:
            with self.subTest(params=params):
                with self.assertRaises(ValidationError):
                    validate_saved_params(params, 'xlsx')
                ReportJob.objects.create(user=self.fed.user, params=params, format='xlsx')
        self.assertEqual(process_reports(limit=100), 0)
        self.assertEqual(ReportJob.objects.filter(status='failed').count(), len(invalid))
        self.assertFalse(report_directory().exists() and list(report_directory().iterdir()))

    def test_changed_source_fails_instead_of_issuing_outdated_preview(self):
        self.competition.protocol_published = False
        self.competition.state = 'closed'
        self.competition.save(update_fields=['protocol_published', 'state'])
        job = self.queue('protocol', self.competition.pk)
        self.competition.state = 'open'
        self.competition.save(update_fields=['state'])
        self.assertEqual(process_reports(), 0)
        job.refresh_from_db()
        self.assertEqual(job.status, 'failed')
        self.assertFalse(job.filename)

    def test_expired_lease_is_recovered_but_current_claim_is_not(self):
        job = self.queue('athlete', self.athlete.pk)
        ReportJob.objects.filter(pk=job.pk).update(status='running', claim_token=uuid.uuid4(), lease_expires_at=timezone.now() + timedelta(minutes=1))
        self.assertEqual(process_reports(), 0)
        ReportJob.objects.filter(pk=job.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(process_reports(), 1)
        self.assertTrue(self.downloaded(job).startswith(b'PK'))

    def test_stale_worker_cannot_overwrite_file_of_newer_claim(self):
        job = self.queue('athlete', self.athlete.pk)
        newer_token = uuid.uuid4()

        def newer_worker_finished(old_job):
            ReportJob.objects.filter(pk=old_job.pk).update(status='ready', claim_token=newer_token, lease_expires_at=None)
            fresh_job = ReportJob.objects.get(pk=old_job.pk)
            filename = expected_filename(fresh_job)
            (report_directory() / filename).write_bytes(b'newer-report')
            ReportJob.objects.filter(pk=fresh_job.pk).update(filename=filename)
            return b'outdated-report'

        with patch('reporting.services.build_report', side_effect=newer_worker_finished):
            self.assertEqual(process_reports(), 0)
        job.refresh_from_db()
        self.assertEqual(job.claim_token, newer_token)
        self.assertEqual(self.downloaded(job), b'newer-report')
        self.assertEqual(len(list(report_directory().iterdir())), 1)

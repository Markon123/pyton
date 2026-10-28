from datetime import date, timedelta
from io import BytesIO, StringIO

from docx import Document
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .forms import CompetitionForm
from .models import (Application, CoachLink, Competition, DictionaryEntry, Discipline,
                     Notification, Organization, Profile, RankApplication, ReportDeadline, Result)


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class CoreUiFixTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.today = timezone.localdate()
        cls.discipline = Discipline.objects.create(name='Алгоритмическое программирование', code='audit-algorithm')
        cls.org = Organization.objects.create(name='Центр «Код Оренбуржья»', municipality='Оренбург')
        cls.fed = cls.person('audit-fed', ['federation'])
        cls.athlete = cls.person('audit-athlete', ['athlete'], 'Александр', 'Тестов')
        cls.other = cls.person('audit-other', ['athlete'], 'Антон', 'Проверочный')
        cls.judge = cls.person('audit-judge', ['judge'])
        cls.coach = cls.person('audit-coach', ['coach'])

    @classmethod
    def person(cls, username, roles, first='Тестовый', last='Участник'):
        user = get_user_model().objects.create_user(username, email=username+'@example.test', password='Test!Strong726Alpha', first_name=first, last_name=last)
        profile = Profile.objects.create(user=user, roles=roles, birth_date=date(cls.today.year-20, 1, 1),
                                         organization=cls.org, municipality='Оренбург', start_year=cls.today.year-3)
        profile.disciplines.add(cls.discipline)
        return profile

    def use(self, person):
        self.client.force_login(person.user)

    def competition(self, **kwargs):
        values = {'title': 'Новый проверочный турнир', 'state': 'closed', 'starts_on': self.today, 'ends_on': self.today,
                  'registration_opens': self.today-timedelta(days=7), 'registration_closes': self.today,
                  'location': 'Проверочная площадка', 'plan_number': 'TEST-001', 'calendar_plan': 'regional',
                  'regulation': 'regulations/test.pdf'}
        values.update(kwargs)
        comp = Competition.objects.create(**values)
        comp.disciplines.add(self.discipline)
        return comp

    def result(self, comp, athlete=None, place=1):
        person = athlete or self.athlete
        application = Application.objects.create(competition=comp, athlete=person, discipline=self.discipline, submitted_by=person.user, status='approved')
        return Result.objects.create(application=application, place=place, score='142.50')

    def test_unicode_fullname_and_organization_search_find_real_result_rows(self):
        self.use(self.fed)
        for query in ('александр', 'АЛЕКСАНДР ТЕСТОВ', 'Тестов Александр', 'код оренбуржья'):
            with self.subTest(query=query):
                response = self.client.get(reverse('registry'), {'q': query})
                self.assertEqual(response.status_code, 200)
                self.assertIn(self.athlete.pk, [person.pk for person in response.context['profiles']])
        self.assertEqual(list(self.client.get(reverse('registry'), {'q': 'Несуществующий Человек'}).context['profiles']), [])

    def test_judge_dashboard_displays_existing_category_and_expiration(self):
        category = RankApplication.objects.create(athlete=self.judge, submitted_by=self.judge.user, kind='category',
            requested_title='Судья второй категории', status='approved', assigned_on=self.today,
            valid_until=self.today+timedelta(days=365), order_number='JUDGE-001')
        self.use(self.judge)
        response = self.client.get(reverse('dashboard'))
        self.assertContains(response, category.requested_title)
        self.assertContains(response, 'JUDGE-001')
        self.assertNotContains(response, 'Действующая категория пока не зарегистрирована')

    def test_calendar_form_allows_draft_but_rejects_wrong_published_plan_and_missing_number(self):
        payload = {'title': 'Проверка календарного плана', 'disciplines': [self.discipline.pk], 'starts_on': self.today,
            'ends_on': self.today, 'registration_opens': self.today, 'registration_closes': self.today,
            'location': 'Тест', 'format': 'online', 'level': 'interregional', 'kind': 'Кубок',
            'age_min': 12, 'age_max': 99, 'calendar_plan': 'regional', 'plan_number': '', 'state': 'draft'}
        self.assertTrue(CompetitionForm(payload).is_valid())
        uploaded = SimpleUploadedFile('regulation.pdf', b'%PDF-1.4\nTest regulation')
        form = CompetitionForm({**payload, 'state': 'published'}, {'regulation': uploaded})
        self.assertFalse(form.is_valid())
        self.assertIn('calendar_plan', form.errors)
        self.assertIn('plan_number', form.errors)
        valid = CompetitionForm({**payload, 'state': 'published', 'calendar_plan': 'ekp', 'plan_number': 'EKP-TEST'},
                                {'regulation': SimpleUploadedFile('regulation.pdf', b'%PDF-1.4\nTest')})
        self.assertTrue(valid.is_valid(), valid.errors)

    def test_existing_invalid_plan_cannot_bypass_checks_via_open_action(self):
        comp = self.competition(state='draft', level='interregional', calendar_plan='regional', plan_number='')
        self.use(self.fed)
        response = self.client.post(reverse('competition_action', args=[comp.pk]), {'action': 'open'})
        self.assertIn('X-Action-Error', response)
        comp.refresh_from_db()
        self.assertEqual(comp.state, 'draft')

    def test_preliminary_protocol_is_real_document_private_and_does_not_publish(self):
        comp = self.competition()
        self.result(comp)
        url = reverse('export_document', args=['protocol', comp.pk, 'docx'])
        self.use(self.fed)
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        document = Document(BytesIO(response.content))
        self.assertIn('Предварительный протокол', '\n'.join(p.text for p in document.paragraphs))
        self.assertIn('Александр Тестов', '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells))
        comp.refresh_from_db()
        self.assertFalse(comp.protocol_published)
        self.use(self.athlete)
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_preliminary_protocol_rejects_incomplete_results(self):
        comp = self.competition()
        Application.objects.create(competition=comp, athlete=self.athlete, discipline=self.discipline, submitted_by=self.athlete.user, status='approved')
        self.use(self.fed)
        self.assertEqual(self.client.get(reverse('export_document', args=['protocol', comp.pk, 'pdf'])).status_code, 400)

    def test_federation_dashboard_period_has_required_counts(self):
        finished = self.competition(state='finished', protocol_published=True)
        result = self.result(finished)
        self.result(finished, self.other, 2)
        self.competition(title='Запланированный старт', state='published')
        RankApplication.objects.create(athlete=self.athlete, submitted_by=self.athlete.user, result=result,
            requested_title='III разряд', status='approved', assigned_on=self.today, valid_until=self.today+timedelta(days=365), order_number='RANK-001')
        RankApplication.objects.create(athlete=self.judge, submitted_by=self.judge.user, kind='category',
            requested_title='Судья второй категории', status='approved', assigned_on=self.today,
            valid_until=self.today+timedelta(days=365), order_number='JUDGE-001')
        self.use(self.fed)
        response = self.client.get(reverse('dashboard'), {'date_from': self.today.isoformat(), 'date_to': self.today.isoformat()})
        stats = {row['label']: row['value'] for row in response.context['stats']}
        self.assertEqual(stats['Проведённых соревнований'], 1)
        self.assertEqual(stats['Запланированных соревнований'], 1)
        self.assertEqual(stats['Участников соревнований'], 2)
        self.assertEqual(stats['Присвоенных разрядов'], 1)
        self.assertEqual(stats['Действующих судейских категорий'], 1)

    def test_history_links_directly_to_published_protocol_and_custom_rank_is_selectable(self):
        comp = self.competition(state='finished', protocol_published=True)
        self.result(comp)
        self.use(self.athlete)
        self.assertContains(self.client.get(reverse('athlete_detail', args=[self.athlete.pk])), reverse('export_document', args=['protocol', comp.pk, 'pdf']))
        DictionaryEntry.objects.create(kind='rank', code='custom-rank', label='Настроенный разряд', metadata={'order': 0})
        self.use(self.fed)
        self.assertContains(self.client.get(reverse('rating')), 'value="Настроенный разряд"')

    def test_report_deadline_reminder_repeats_safely_and_stops_after_completion(self):
        deadline = ReportDeadline.objects.create(title='Годовой отчёт', due_on=self.today+timedelta(days=2), remind_days=7, created_by=self.fed.user)
        call_command('send_reminders', stdout=StringIO())
        notices = Notification.objects.filter(title__contains=f'Отчёт № {deadline.pk}')
        self.assertEqual(notices.count(), 1)
        self.assertEqual(notices.get().user, self.fed.user)
        call_command('send_reminders', stdout=StringIO())
        self.assertEqual(notices.count(), 1)
        deadline.completed_at = timezone.now()
        deadline.save()
        notices.delete()
        call_command('send_reminders', stdout=StringIO())
        self.assertEqual(notices.count(), 0)

    def test_closing_registration_notifies_eligible_athlete_and_coach_without_application(self):
        self.competition(state='open', registration_closes=self.today+timedelta(days=1), starts_on=self.today+timedelta(days=2), ends_on=self.today+timedelta(days=2))
        CoachLink.objects.create(coach=self.coach, athlete=self.athlete, discipline=self.discipline, status='active')
        self.assertFalse(Application.objects.exists())
        call_command('send_reminders', stdout=StringIO())
        self.assertTrue(Notification.objects.filter(user=self.athlete.user).exists())
        self.assertTrue(Notification.objects.filter(user=self.coach.user).exists())
        self.assertTrue(Notification.objects.filter(user=self.fed.user).exists())

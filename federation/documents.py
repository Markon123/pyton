"""Neutral demonstration exports; authorization belongs to their calling view."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
from xml.sax.saxutils import escape

from django.http import Http404, HttpResponse
from django.utils import timezone
from django.utils.http import content_disposition_header


DISCLAIMER = "Демонстрационный документ · вымышленные данные"


def document_note():
    from django.conf import settings
    return DISCLAIMER if settings.FSP_DEMO_MODE else 'Сформировано в платформе ОренКод'
CONTENT_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%d.%m.%Y")
    # Control characters in user text must not break XML-based document formats.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def _cell(value):
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return value
    value = _text(value)
    # Also cover leading whitespace, which spreadsheet software may ignore.
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value[:32767]


def _xlsx(title, headers, rows, subtitle, generated):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Данные"
    sheet.append([_cell(title)])
    sheet.append([_cell(subtitle)])
    sheet.append([_cell(f"Сформирован: {generated} · {document_note()}")])
    sheet.append([])
    sheet.append([_cell(header) for header in headers])
    for row in rows:
        sheet.append([_cell(value) for value in row])
    for cell in sheet[1]:
        cell.font = Font(size=16, bold=True, color="152137")
    for cell in sheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="152137")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=6):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index in range(1, len(headers) + 1):
        values = [sheet.cell(row, index).value for row in range(5, sheet.max_row + 1)]
        width = min(55, max(16, max((len(str(value or "")) for value in values), default=16) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{get_column_letter(len(headers))}{max(5, sheet.max_row)}"
    sheet.print_title_rows = "1:5"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _docx(title, headers, rows, subtitle, generated):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.shared import Cm, Pt

    document = Document()
    section = document.sections[0]
    if len(headers) > 5:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.5)
    document.styles["Normal"].font.name = "DejaVu Sans"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading(title, level=0)
    if subtitle:
        document.add_paragraph(subtitle)
    document.add_paragraph(f"Сформирован: {generated}")
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = _text(value)
    section.footer.paragraphs[0].text = document_note()
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf(title, headers, rows, subtitle, generated):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font_directory = Path(__file__).parent / "assets" / "fonts"
    if "FspDejaVu" not in pdfmetrics.getRegisteredFontNames():
        for name, filename in (("FspDejaVu", "DejaVuSans.ttf"), ("FspDejaVu-Bold", "DejaVuSans-Bold.ttf")):
            bundled = font_directory / filename
            source = bundled if bundled.is_file() else Path("/usr/share/fonts/truetype/dejavu") / filename
            pdfmetrics.registerFont(TTFont(name, str(source)))
        pdfmetrics.registerFontFamily("FspDejaVu", normal="FspDejaVu", bold="FspDejaVu-Bold")
    pagesize = landscape(A4) if len(headers) > 5 else A4
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=pagesize, rightMargin=28, leftMargin=28,
                                 topMargin=32, bottomMargin=46, title=title, author="ОренКод")
    base = ParagraphStyle("FspBody", fontName="FspDejaVu", fontSize=8, leading=12,
                          alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle("FspHeading", parent=base, fontName="FspDejaVu-Bold", fontSize=16, leading=21)
    header = ParagraphStyle("FspHeader", parent=base, fontName="FspDejaVu-Bold", textColor=colors.white)

    def paragraph(value, style=base):
        return Paragraph(escape(_text(value)).replace("\n", "<br/>"), style)

    story = [paragraph(title, heading), Spacer(1, 12)]
    if subtitle:
        story.extend([paragraph(subtitle), Spacer(1, 8)])
    story.extend([paragraph(f"Сформирован: {generated}"), Spacer(1, 14)])
    data = [[paragraph(value, header) for value in headers]]
    data.extend([[paragraph(value) for value in row] for row in rows])
    table = LongTable(data, colWidths=[document.width / len(headers)] * len(headers),
                      repeatRows=1, splitInRow=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#152137")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f3f5")]),
        ("LINEBELOW", (0, 0), (-1, 0), 0.4, colors.HexColor("#cad2db")),
    ]))
    story.append(table)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("FspDejaVu", 7)
        canvas.setFillColor(colors.HexColor("#596579"))
        canvas.drawString(28, 25, document_note())
        canvas.drawRightString(pagesize[0] - 28, 25, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def build_document(kind, fmt, title, headers, rows, subtitle=""):
    """Export tabular data. The caller must authorize every supplied row."""
    if fmt not in CONTENT_TYPES:
        raise Http404("Неизвестный формат документа")
    headers = [_text(value) for value in headers]
    if not headers:
        raise ValueError("Document requires at least one column")
    rows = [list(row) for row in rows]
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("Every document row must match its headers")
    generated = timezone.localtime().strftime("%d.%m.%Y %H:%M")
    data = {"xlsx": _xlsx, "docx": _docx, "pdf": _pdf}[fmt](
        _text(title), headers, rows, _text(subtitle), generated)
    filename = f"{re.sub(r'[^\w.-]', '_', str(kind)) or 'document'}.{fmt}"
    response = HttpResponse(data, content_type=CONTENT_TYPES[fmt])
    response["Content-Disposition"] = content_disposition_header(True, filename)
    response["Cache-Control"] = "private, no-store"
    return response

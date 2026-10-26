"""Repeatable fictitious fixtures. Existing objects and user edits are preserved."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from docx import Document

from django.contrib.auth import get_user_model
from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone

from federation.models import (
    Application, AuditLog, CoachLink, Competition, Consent, Discipline,
    JudgeAssignment, Notification, Organization, Profile, RankApplication,
    RatingRule, Result, Role, Team,
)


class Command(BaseCommand):
    help = 'Добавить вымышленные демонстрационные данные без изменения существующих записей.'

    @transaction.atomic
    def handle(self, *args, **options):
        today = timezone.localdate()
        year = today.year
        password = 'FspDemo2026!'
        User = get_user_model()
        disciplines = []
        for code, name in [
            ('algorithm', 'Программирование алгоритмическое'),
            ('product', 'Программирование продуктовое'),
            ('security', 'Системы информационной безопасности'),
            ('ai', 'Системы искусственного интеллекта'),
            ('robotics', 'Робототехника'),
            ('drone', 'Беспилотные авиационные системы'),
        ]:
            obj, _ = Discipline.objects.get_or_create(code=code, defaults={'name': name})
            disciplines.append(obj)
        organizations = []
        for name, kind, city in [
            ('Демоцентр «Код Оренбуржья»', 'Спортивный клуб', 'Оренбург'),
            ('Демошкола цифрового спорта', 'Спортивная школа', 'Орск'),
            ('Демоакадемия «Импульс»', 'Образовательный центр', 'Бузулук'),
            ('Демоклуб «Точка старта»', 'Спортивный клуб', 'Новотроицк'),
        ]:
            obj, _ = Organization.objects.get_or_create(name=name, defaults={'kind': kind, 'municipality': city})
            organizations.append(obj)

        def make_user(username, first, last, roles, index=0):
            user, created = User.objects.get_or_create(username=username, defaults={
                'first_name': first, 'last_name': last, 'email': f'{username}@example.test',
            })
            if created:
                user.set_password(password)
                user.save()
            age = 17 + index % 10 if Role.ATHLETE in roles else 32 + index
            birthday = date(year - age, 4, 10)
            profile, profile_created = Profile.objects.get_or_create(user=user, defaults={
                'roles': roles, 'birth_date': birthday,
                'municipality': organizations[index % 4].municipality,
                'organization': organizations[index % 4], 'start_year': year - 4 + index % 4,
                'guardian_name': 'Демо Представитель Алексей Сергеевич' if (today.year - birthday.year - ((today.month, today.day) < (birthday.month, birthday.day))) < 18 else '',
                'anti_doping_date': today - timedelta(days=60),
                'anti_doping_url': 'https://rusada.ru/education/',
            })
            if profile_created:
                profile.disciplines.set([disciplines[index % 6], disciplines[(index + 1) % 6]])
            if created or profile_created:
                Consent.objects.get_or_create(user=user, kind='personal_data', version='1.0')
                if profile.age is not None and profile.age < 18:
                    Consent.objects.get_or_create(user=user, kind='guardian', version='1.0', defaults={'representative_name': profile.guardian_name})
            return profile

        federation = make_user('federation', 'Мария', 'Демонстрационная', list(Role.values), 1)
        athlete = make_user('athlete', 'Александр', 'Тестов', [Role.ATHLETE], 0)
        coach = make_user('coach', 'Ирина', 'Учебная', [Role.COACH], 2)
        judge = make_user('judge', 'Михаил', 'Примеров', [Role.JUDGE], 3)
        athletes = [athlete]
        names = [
            ('Алина', 'Пробная'), ('Даниил', 'Макетный'), ('Софья', 'Тестовая'),
            ('Артём', 'Примерный'), ('Виктория', 'Учебная'), ('Илья', 'Демонстрационный'),
            ('Полина', 'Образцовая'), ('Максим', 'Проверочный'), ('Дарья', 'Макетная'),
            ('Никита', 'Тренировочный'), ('Елизавета', 'Примерова'), ('Кирилл', 'Демов'),
            ('Анастасия', 'Учебнова'), ('Матвей', 'Тестовский'), ('Вероника', 'Пробникова'),
            ('Егор', 'Образцовский'), ('Ксения', 'Проверочная'), ('Роман', 'Демо'),
            ('Мария', 'Сценарная'), ('Лев', 'Пилотный'), ('Варвара', 'Макетова'),
            ('Тимофей', 'Учебников'), ('Анна', 'Синтетическая'),
        ]
        for index, (first, last) in enumerate(names, start=1):
            athletes.append(make_user(f'demo_athlete_{index:02}', first, last, [Role.ATHLETE], index))
        for profile in athletes[:9]:
            for discipline in profile.disciplines.all():
                CoachLink.objects.get_or_create(coach=coach, athlete=profile, discipline=discipline, defaults={'status': 'active'})

        for level, points in [
            ('regional', (40, 30, 20)), ('interregional', (60, 45, 30)),
            ('russian_championship', (100, 80, 60)), ('russian_junior', (80, 60, 40)),
        ]:
            RatingRule.objects.get_or_create(level=level, defaults=dict(zip(['first', 'second', 'third'], points)))

        completed = []
        year_start = date(year, 1, 1)
        elapsed = max(0, (today - year_start).days)
        for index, (title, level) in enumerate([
            ('Зимний кубок цифрового спорта', 'regional'),
            ('Открытый чемпионат Оренбуржья', 'regional'),
            ('Кубок Приволжья', 'interregional'),
        ]):
            starts = year_start + timedelta(days=int(elapsed * (index + 1) / 5))
            obj, created = Competition.objects.get_or_create(title=f'{title} · демо {year}', defaults={
                'description': 'Демонстрационное соревнование. Все участники и результаты вымышлены.',
                'starts_on': starts, 'ends_on': starts,
                'registration_opens': starts - timedelta(days=30), 'registration_closes': starts - timedelta(days=2),
                'location': 'Демоцентр «Код Оренбуржья», Оренбург', 'format': 'hybrid',
                'level': level, 'kind': 'Кубок' if index != 1 else 'Чемпионат',
                'age_min': 12, 'age_max': 99, 'plan_number': f'ДЕМО-{year}-{index + 1:03}',
                'state': 'finished', 'protocol_published': True,
            })
            if created:
                obj.disciplines.set(disciplines)
                self.regulation(obj)
                JudgeAssignment.objects.get_or_create(competition=obj, judge=judge, defaults={'position': 'Главный судья'})
                for d_index, discipline in enumerate(disciplines):
                    group = [p for number, p in enumerate(athletes) if number % 6 == d_index]
                    for place, profile in enumerate(group, start=1):
                        app = Application.objects.create(competition=obj, athlete=profile, discipline=discipline, submitted_by=profile.user, status='approved')
                        Result.objects.create(application=app, place=place, score=Decimal(1000 - place * 83 + index * 12))
            completed.append(obj)

        upcoming = []
        for index, title in enumerate([
            'Чемпионат Оренбургской области', 'Хакатон «ОренКод»', 'Кубок по кибербезопасности',
            'Турнир «Интеллект в коде»', 'Фестиваль робототехники и БАС',
        ]):
            starts = today + timedelta(days=12 + index * 9)
            obj, created = Competition.objects.get_or_create(title=f'{title} · демо {year}', defaults={
                'description': 'Присоединяйтесь к соревнованию по спортивному программированию. '
                'Регистрация доступна спортсменам всех организаций региона. Демонстрационные данные.',
                'starts_on': starts, 'ends_on': starts + timedelta(days=1 if index == 1 else 0),
                'registration_opens': today - timedelta(days=8), 'registration_closes': starts - timedelta(days=3),
                'location': ['Оренбург · Демоцентр «Код Оренбуржья»', 'Онлайн · демоплощадка', 'Орск · Демошкола цифрового спорта'][index % 3],
                'format': ['offline', 'hybrid', 'online'][index % 3],
                'level': 'interregional' if index == 1 else 'regional', 'kind': 'Чемпионат' if index == 0 else 'Кубок',
                'age_min': 12, 'age_max': 99, 'plan_number': f'ДЕМО-{year}-{index + 10:03}', 'state': 'open',
            })
            if created:
                obj.disciplines.set([disciplines[index]] if index < 4 else disciplines[4:])
                self.regulation(obj)
                JudgeAssignment.objects.get_or_create(competition=obj, judge=judge, defaults={'position': 'Главный судья'})
                if index == 0:
                    for profile in athletes[:5]:
                        Application.objects.get_or_create(competition=obj, athlete=profile, discipline=disciplines[0], defaults={
                            'submitted_by': profile.user, 'status': 'pending',
                        })
            upcoming.append(obj)

        main_result = Result.objects.filter(application__athlete=athlete, application__competition=completed[0]).first()
        if main_result:
            RankApplication.objects.get_or_create(athlete=athlete, requested_title='II разряд', kind='rank', result=main_result, defaults={
                'submitted_by': federation.user, 'status': 'approved', 'assigned_on': main_result.application.competition.starts_on + timedelta(days=1),
                'valid_until': today + timedelta(days=25), 'order_number': f'ДЕМО-{year}-014',
                'review_comment': 'Вымышленный приказ для демонстрации контроля сроков.',
            })
            RankApplication.objects.get_or_create(athlete=athlete, requested_title='I разряд', kind='rank', result=main_result, defaults={
                'submitted_by': athlete.user, 'status': 'pending',
            })
        RankApplication.objects.get_or_create(athlete=judge, kind='category', requested_title='Судья второй категории', defaults={
            'submitted_by': federation.user, 'status': 'approved', 'assigned_on': today - timedelta(days=340),
            'valid_until': today + timedelta(days=20), 'order_number': f'ДЕМО-{year}-С02',
        })
        RankApplication.objects.get_or_create(athlete=judge, kind='category', requested_title='Судья первой категории', defaults={
            'submitted_by': judge.user, 'status': 'pending',
        })
        for profile, title, body, url in [
            (federation, 'Заявки ждут решения', 'Откройте чемпионат области и рассмотрите заявки участников.', f'/competitions/{upcoming[0].pk}/'),
            (athlete, 'Добро пожаловать в ФСП', 'Ваша история результатов и сроки разрядов — в личном кабинете.', '/profile/'),
            (athlete, 'Срок разряда подходит к концу', 'До окончания действия II разряда осталось 25 дней.', '/ranks/'),
            (coach, 'Спортсмены привязаны к тренеру', 'Откройте кабинет тренера, чтобы подать заявки за спортсменов.', '/coaching/'),
            (judge, 'Новое судейское назначение', 'Вы назначены главным судьёй чемпионата области.', f'/competitions/{upcoming[0].pk}/'),
        ]:
            Notification.objects.get_or_create(user=profile.user, title=title, defaults={'body': body, 'url': url})
        AuditLog.objects.get_or_create(user=federation.user, role=Role.FEDERATION, action='Загружены демонстрационные данные', entity='seed_demo', defaults={
            'detail': 'Все профили, организации, результаты и приказы вымышлены.',
        })
        self.stdout.write(self.style.SUCCESS(
            f'Демо-данные готовы: {Discipline.objects.count()} дисциплин, '
            f'{Organization.objects.count()} организаций, {Profile.objects.count()} профилей, '
            f'{Competition.objects.count()} соревнований. Существующие данные сохранены.'
        ))
        self.stdout.write('Демо-вход: federation / athlete / coach / judge; пароль: FspDemo2026!')

    @staticmethod
    def regulation(competition):
        document = Document()
        document.add_heading('ПОЛОЖЕНИЕ О СОРЕВНОВАНИИ', 0)
        document.add_paragraph('ДЕМОНСТРАЦИОННЫЙ ДОКУМЕНТ. Все сведения вымышлены.')
        document.add_heading(competition.title, level=1)
        document.add_paragraph(f'Дата: {competition.starts_on:%d.%m.%Y}')
        document.add_paragraph(f'Место: {competition.location}')
        document.add_paragraph(f'Возраст участников: {competition.age_min}–{competition.age_max} лет.')
        document.add_paragraph(f'Регистрация до {competition.registration_closes:%d.%m.%Y}.')
        document.add_heading('Порядок участия', level=2)
        document.add_paragraph('Участие доступно после подачи и одобрения заявки федерацией. '
                               'Победители определяются по результатам соревнования отдельно в каждой дисциплине.')
        document.add_paragraph('Настоящий файл создан для демонстрации платформы и не является официальным положением.')
        content = BytesIO()
        document.save(content)
        competition.regulation.save(f'demo-{competition.pk}.docx', ContentFile(content.getvalue()), save=True)


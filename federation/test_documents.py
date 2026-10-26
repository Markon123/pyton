from io import BytesIO
from zipfile import ZipFile

from django.http import Http404
from django.test import SimpleTestCase
from docx import Document
from openpyxl import load_workbook

from .documents import DISCLAIMER, build_document


class DocumentExportTests(SimpleTestCase):
    def export(self, fmt, rows=None):
        return build_document("протокол", fmt, "Протокол соревнования", ["Место", "Спортсмен", "Баллы"],
                              rows if rows is not None else [[1, "Тестовый Спортсмен", 45]], "Алгоритмическое программирование")

    def test_xlsx_is_readable_and_has_real_numeric_cells(self):
        response = self.export("xlsx")
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["B6"].value, "Тестовый Спортсмен")
        self.assertEqual(sheet["C6"].value, 45)
        self.assertEqual(sheet["C6"].data_type, "n")
        self.assertEqual(sheet.freeze_panes, "A6")
        self.assertEqual(sheet.auto_filter.ref, "A5:C6")
        self.assertIn("filename*=utf-8", response["Content-Disposition"])

    def test_xlsx_user_formulas_are_not_executable(self):
        response = self.export("xlsx", [[1, '=HYPERLINK("https://example.test")', " @SUM(A1:A2)"],
                                        [2, "+cmd|' /C calc'!A0", "-2+3"]])
        sheet = load_workbook(BytesIO(response.content)).active
        for address in ["B6", "C6", "B7", "C7"]:
            self.assertEqual(sheet[address].data_type, "s")
            self.assertTrue(sheet[address].value.startswith("'"))

    def test_docx_has_document_table_and_demo_disclaimer(self):
        response = self.export("docx")
        document = Document(BytesIO(response.content))
        self.assertEqual(document.tables[0].rows[1].cells[1].text, "Тестовый Спортсмен")
        self.assertIn(DISCLAIMER, document.sections[0].footer.paragraphs[0].text)
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertIn(b"tblHeader", archive.read("word/document.xml"))

    def test_pdf_has_embedded_cyrillic_font_and_valid_header(self):
        response = self.export("pdf", [[1, "Тестовый <спортсмен> & команда", 45]])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(b"DejaVuSans", response.content)
        self.assertIn(b"/FontFile2", response.content)
        self.assertIn(b"%%EOF", response.content)

    def test_pdf_long_cells_split_across_pages(self):
        response = self.export("pdf", [[1, ("Очень длинное имя и описание " * 900), 45]])
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertGreater(len(response.content), 10000)

    def test_empty_exports_and_control_characters_are_safe(self):
        for fmt in ("xlsx", "docx", "pdf"):
            self.assertEqual(self.export(fmt, []).status_code, 200)
            self.assertEqual(self.export(fmt, [[1, "Тест\x00овый", 10]]).status_code, 200)

    def test_unsupported_format_is_a_404(self):
        with self.assertRaises(Http404):
            self.export("exe")

    def test_mismatched_row_rejected(self):
        with self.assertRaises(ValueError):
            self.export("pdf", [[1, "неполная строка"]])
